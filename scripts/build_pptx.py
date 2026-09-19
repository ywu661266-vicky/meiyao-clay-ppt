#!/usr/bin/env python3
"""Meiyao Clay PPT —— 背景图 → 16:9 PPTX 合成脚本。

把 ImageGen 生成的黏土风页面铺满 16:9 幻灯片，并用**真实字体**叠加正文文字。

【核心原则（解决"字乱码/写不进内容"）】
    图片模型画不好中文——把成段文字交给它"捏成黏土字"，必然错字乱码，
    而且一页只能塞下标题级字数。所以正文一律**不烘进图**，
    改由本脚本用系统字体（微软雅黑等）真实叠字，清晰、可编辑、字数不受限。

【text_mode="baked"】
    图里的标题/标签已经是「黏土立体字」（烘进画面），脚本不再叠加任何扁平文字。

【text_mode="overlay"】（内容页推荐）
    在图中预留的空白区，用真字体叠加正文。支持三种版式：
      layout="cards"   → 卡片列表（head 关键词 + body 说明），内容页首选
      layout="steps"   → 步骤条（序号 + 标题 + 说明）
      layout="bullets" → 纯项目符号列表（兼容旧版）
    content_side 决定文字落在哪一侧（left / right / full）。

用法：
    python build_pptx.py --config slides.json --out "主题_黏土风.pptx"
    python build_pptx.py --config slides.json --out out.pptx --font "Microsoft YaHei"

slides.json 结构：
{
  "title": "整套课的标题（仅用于元数据）",
  "default_font": "Microsoft YaHei",
  "default_text_mode": "baked",              // baked（默认）| overlay
  "slides": [
    { "image": "cover.png", "text_mode": "baked", "notes": "讲师备注" },
    {
      "image": "content.png",
      "text_mode": "overlay",
      "layout": "cards",                     // cards | steps | bullets
      "content_side": "left",                // left | right | full
      "items": [                             // layout=cards / steps 时用
        { "head": "关键词", "body": "一句话说明" },
        { "head": "关键词", "body": "一句话说明" }
      ],
      "bullets": ["要点一", "要点二"],         // layout=bullets 时用
      "notes": "讲师备注",
      "accent_color": "6B4BB5",              // 可选，卡片左沿强调色
      "chead_color": "4A2C6B",               // 可选，关键词颜色
      "cbody_color": "3A3350"                // 可选，说明文字颜色
    }
  ]
}
"""
from __future__ import annotations

import argparse
import json
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
EMU_PER_INCH = 914400

DEFAULT_FONT = "Microsoft YaHei"
TITLE_COLOR_DEFAULT = "1E3A5F"
SUBTITLE_COLOR = "3E5C7A"
BODY_COLOR = "24384F"
PANEL_FILL = "FFF6E3"

# —— 内容页叠字版式默认色（深紫卡通风，和黏土紫底协调）——
CARD_FILL = "FFF6E3"          # 卡片奶油底
CARD_ACCENT = "6B4BB5"        # 卡片左沿强调色
CARD_HEAD = "4A2C6B"          # 关键词
CARD_BODY = "3A3350"          # 说明


def hex_rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def set_run_font(run, name: str, size: float, bold: bool = False, color: RGBColor | None = None) -> None:
    """设置字体，并同时写入 latin / ea / cs，保证中文不回退成宋体。"""
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    if color is not None:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def set_alpha(shape, percent: int) -> None:
    """给纯色填充加透明度（percent = 不透明度 0-100）。"""
    solid = shape.fill._xPr.find(qn("a:solidFill"))
    if solid is None:
        return
    clr = solid.find(qn("a:srgbClr"))
    if clr is None:
        return
    for child in list(clr):
        if child.tag == qn("a:alpha"):
            clr.remove(child)
    alpha = clr.makeelement(qn("a:alpha"), {"val": str(int(percent * 1000))})
    clr.append(alpha)


def add_panel(slide, left: Emu, top: Emu, width: Emu, height: Emu):
    """奶油色半透明圆角底板，保证文字可读。"""
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = hex_rgb(PANEL_FILL)
    set_alpha(box, 88)
    box.line.fill.background()
    box.shadow.inherit = False
    return box


def add_textbox(slide, left, top, width, height, anchor=MSO_ANCHOR.MIDDLE):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tf


def _fill_rounded(shape, color_hex: str, alpha: int | None = None) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = hex_rgb(color_hex)
    if alpha is not None:
        set_alpha(shape, alpha)
    shape.line.fill.background()
    shape.shadow.inherit = False


def _content_zone(side: str):
    """返回 (left, top, width, height)，即预留的叠字区（EMU）。"""
    top, bottom = Inches(1.95), Inches(6.95)
    h = bottom - top
    if side == "right":
        return Inches(8.55), top, Inches(4.05), h
    if side == "full":
        return Inches(0.9), top, SLIDE_W - Inches(1.8), h
    # left（默认）
    return Inches(0.72), top, Inches(6.95), h


def render_cards(slide, zone, items, font, accent, chead, cbody):
    left, top, width, height = zone
    n = len(items)
    if n == 0:
        return
    gap = Inches(0.22)
    card_h = int((height - gap * (n - 1)) / n)
    for i, it in enumerate(items):
        y = top + i * (card_h + gap)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, y, width, card_h)
        _fill_rounded(card, CARD_FILL, 92)
        # 左沿强调条
        bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, y + Inches(0.12),
                                     Inches(0.13), card_h - Inches(0.24))
        _fill_rounded(bar, accent)
        # 序号徽章
        badge_d = Inches(0.56)
        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, left + Inches(0.32), y + int((card_h - badge_d) / 2), badge_d, badge_d)
        _fill_rounded(badge, accent)
        btf = badge.text_frame
        btf.word_wrap = False
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE
        bp = btf.paragraphs[0]
        bp.alignment = PP_ALIGN.CENTER
        set_run_font(bp.add_run(), font, 16, True, hex_rgb("FFFFFF"))
        bp.runs[0].text = str(i + 1)
        # 文字
        tx = left + Inches(1.12)
        tw = width - Inches(1.42)
        tf = add_textbox(slide, tx, y + Inches(0.18), tw, card_h - Inches(0.36),
                         anchor=MSO_ANCHOR.MIDDLE)
        head = str(it.get("head", "") or "")
        body = str(it.get("body", "") or "")
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = 1.05
        if head:
            set_run_font(p.add_run(), font, 17, True, hex_rgb(chead))
            p.runs[-1].text = head
        if body:
            if head:
                p2 = tf.add_paragraph()
                p2.alignment = PP_ALIGN.LEFT
                p2.space_before = Pt(3)
                p2.line_spacing = 1.18
                set_run_font(p2.add_run(), font, 13, False, hex_rgb(cbody))
                p2.runs[-1].text = body
            else:
                set_run_font(p.add_run(), font, 13.5, False, hex_rgb(cbody))
                p.runs[-1].text = body


def render_steps(slide, zone, items, font, accent, chead, cbody):
    """纵向步骤条：圆点 + 竖线 + 标题说明。"""
    left, top, width, height = zone
    n = len(items)
    if n == 0:
        return
    gap = Inches(0.2)
    step_h = int((height - gap * (n - 1)) / n)
    line_x = left + Inches(0.28)
    if n > 1:
        conn = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, line_x, top + Inches(0.3),
                                      Inches(0.035), height - Inches(0.6))
        _fill_rounded(conn, accent, 55)
    d = Inches(0.52)
    for i, it in enumerate(items):
        y = top + i * (step_h + gap)
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, line_x - d // 2 + Inches(0.017), y, d, d)
        _fill_rounded(dot, accent)
        dtf = dot.text_frame
        dtf.word_wrap = False
        dtf.vertical_anchor = MSO_ANCHOR.MIDDLE
        dp = dtf.paragraphs[0]
        dp.alignment = PP_ALIGN.CENTER
        set_run_font(dp.add_run(), font, 15, True, hex_rgb("FFFFFF"))
        dp.runs[0].text = str(i + 1)
        tx = left + Inches(0.78)
        tw = width - Inches(0.95)
        tf = add_textbox(slide, tx, y, tw, step_h, anchor=MSO_ANCHOR.MIDDLE)
        head = str(it.get("head", "") or "")
        body = str(it.get("body", "") or "")
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = 1.05
        if head:
            set_run_font(p.add_run(), font, 17, True, hex_rgb(chead))
            p.runs[-1].text = head
        if body:
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.LEFT
            p2.space_before = Pt(3)
            p2.line_spacing = 1.18
            set_run_font(p2.add_run(), font, 13, False, hex_rgb(cbody))
            p2.runs[-1].text = body


def render_bullets(slide, zone, item, font):
    """兼容旧版：单块奶油底板 + 标题/副标题/项目符号。"""
    left, top, width, height = zone
    title = item.get("title")
    subtitle = item.get("subtitle")
    bullets = item.get("bullets") or []
    if title or bullets or subtitle:
        add_panel(slide, left, top, width, height)
    pad_x = Inches(0.35)
    text_left = left + pad_x
    text_w = width - pad_x * 2
    if title:
        title_h = Inches(1.25) if bullets else height - Inches(0.5)
        tf = add_textbox(slide, text_left, top + Inches(0.22), text_w, title_h)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        set_run_font(p.add_run(), font, 30 if bullets else 36, True,
                     hex_rgb(item.get("title_color", TITLE_COLOR_DEFAULT)))
        p.runs[-1].text = str(title)
    if subtitle:
        tf = add_textbox(slide, text_left, top + Inches(1.35), text_w, Inches(0.6))
        set_run_font(tf.paragraphs[0].add_run(), font, 16, False, hex_rgb(SUBTITLE_COLOR))
        tf.paragraphs[0].runs[-1].text = str(subtitle)
    if bullets:
        body_top = top + (Inches(1.95) if subtitle else Inches(1.6))
        body_h = height - (body_top - top) - Inches(0.3)
        tf = add_textbox(slide, text_left, body_top, text_w, body_h, anchor=MSO_ANCHOR.TOP)
        for i, b in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.space_after = Pt(10)
            p.line_spacing = 1.3
            set_run_font(p.add_run(), font, 16, False, hex_rgb(BODY_COLOR))
            p.runs[-1].text = f"· {b}"


def build_slide(prs: Presentation, item: dict, default_font: str,
                default_text_mode: str = "baked") -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    image = item.get("image")
    if image and os.path.exists(image):
        slide.shapes.add_picture(image, 0, 0, width=prs.slide_width, height=prs.slide_height)
    else:
        print(f"[warn] 背景图不存在，已跳过: {image}", file=sys.stderr)

    notes = item.get("notes")
    if notes:
        slide.notes_slide.notes_text_frame.text = str(notes)

    mode = (item.get("text_mode") or default_text_mode).lower()

    # ---- baked：图里已有黏土立体字，脚本不再叠任何扁平文字 ----
    if mode != "overlay":
        if item.get("title") or item.get("subtitle") or item.get("bullets") or item.get("items"):
            print(
                f"[warn] text_mode=baked：图为 {os.path.basename(str(image))}，"
                "已跳过脚本叠字（避免双层字）。要叠正文请设 text_mode=\"overlay\"。",
                file=sys.stderr,
            )
        return

    # ---- overlay：用真实字体叠正文 ----
    side = (item.get("content_side") or item.get("text_side") or "left").lower()
    if side == "top":
        zone = (Inches(1.4), Inches(0.5), SLIDE_W - Inches(2.8), Inches(1.9))
    else:
        zone = _content_zone(side)

    font = item.get("font") or default_font
    accent = item.get("accent_color", CARD_ACCENT)
    chead = item.get("chead_color", CARD_HEAD)
    cbody = item.get("cbody_color", CARD_BODY)
    layout = (item.get("layout") or "").lower()

    items = item.get("items") or []
    if not layout:
        layout = "cards" if items else "bullets"

    if layout == "cards" and items:
        render_cards(slide, zone, items, font, accent, chead, cbody)
    elif layout == "steps" and items:
        render_steps(slide, zone, items, font, accent, chead, cbody)
    else:
        render_bullets(slide, zone, item, font)


def main() -> int:
    ap = argparse.ArgumentParser(description="把黏土风背景图合成为可编辑 16:9 PPTX")
    ap.add_argument("--config", required=True, help="slides.json 路径")
    ap.add_argument("--out", required=True, help="输出 .pptx 路径")
    ap.add_argument("--font", default=None, help=f"中文字体，默认 {DEFAULT_FONT}")
    args = ap.parse_args()

    with open(args.config, "r", encoding="utf-8") as f:
        cfg = json.load(f)

    font = args.font or cfg.get("default_font") or DEFAULT_FONT
    text_mode = (cfg.get("default_text_mode") or "baked").lower()
    slides = cfg.get("slides") or []
    if not slides:
        print("[error] slides.json 里没有 slides", file=sys.stderr)
        return 1

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    if cfg.get("title"):
        prs.core_properties.title = str(cfg["title"])

    for item in slides:
        build_slide(prs, item, font, text_mode)

    out_dir = os.path.dirname(os.path.abspath(args.out))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    prs.save(args.out)
    print(f"[ok] 已生成 {len(slides)} 页：{args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
